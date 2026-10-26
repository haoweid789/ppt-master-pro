#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_image_deck.py — 整页 PNG 目录 → 图片版 PPTX + PDF(ppt-master-pro 第三输出模式)

路线源自 D-PPT-skill(中文商业汇报图片版工作流):图像模型整页生成 PNG 后,
封装为"每页一张整图"的稳定 PPTX 与 PDF——适合品牌视觉型页面(AI 生图的美学
是原生形状给不了的),代价是不可逐对象编辑(路线定位,非缺陷)。

铁律(与 D-PPT 一致):不 mock 不假图兜底——图像接口失败就明确停止,
缺几页报几页,不用占位图凑数。

用法:
    python build_image_deck.py output/pages --out deck.pptx --pdf deck.pdf
    python build_image_deck.py output/pages --out deck.pptx --expect 12
验证(交付前三查):页数与 expect 一致、每张图宽高比≈16:9、文件可解析。
"""
import argparse
import os
import sys
import glob

try:
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
except ImportError as _e:  # 依赖缺失给出可执行提示
    sys.exit(f"缺少依赖({_e})。安装:{sys.executable} -m pip install python-pptx Pillow")

W, H = 13.333, 7.5
CANVAS_RATIO = W / H
TOL = 0.02  # 宽高比容差 2%


def find_pages(images_dir):
    pages = sorted(
        glob.glob(os.path.join(images_dir, "*.png"))
        + glob.glob(os.path.join(images_dir, "*.jpg"))
        + glob.glob(os.path.join(images_dir, "*.jpeg")),
        key=lambda p: (os.path.getmtime(p), p),
    )
    return pages


def main():
    ap = argparse.ArgumentParser(description="整页 PNG → 图片版 PPTX + PDF")
    ap.add_argument("images_dir", help="逐页图片目录(page-1.png …)")
    ap.add_argument("--out", required=True, help="输出 .pptx 路径")
    ap.add_argument("--pdf", help="同时输出 PDF 路径(可选)")
    ap.add_argument("--expect", type=int, help="期望页数(不匹配即失败)")
    args = ap.parse_args()

    pages = find_pages(args.images_dir)
    if not pages:
        sys.exit(f"FAIL {args.images_dir} 里没有图片——生图阶段没有产出,拒绝封装空壳(不 mock 兜底)")
    if args.expect and len(pages) != args.expect:
        sys.exit(f"FAIL 页数不符: 有 {len(pages)} 页,期望 {args.expect}——"
                 f"断点续跑:补生成缺失页后重新封装")

    # 逐张校验:可解析 + 宽高比
    bad = []
    for p in pages:
        try:
            with Image.open(p) as im:
                w, h = im.size
                ratio = w / h
                if abs(ratio - CANVAS_RATIO) / CANVAS_RATIO > TOL:
                    bad.append(f"  [宽高比偏差] {os.path.basename(p)}: "
                               f"{w}x{h} ({ratio:.3f} vs 画布 {CANVAS_RATIO:.3f})")
        except Exception as e:
            bad.append(f"  [不可解析] {os.path.basename(p)}: {e}")
    if bad:
        print("WARN 图片质量检查发现问题(不阻塞封装,但交付前必须处置):")
        print("\n".join(bad))

    # 封装 PPTX:每页一张全幅图
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for p in pages:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(p, 0, 0, width=Inches(W), height=Inches(H))
    out_dir = os.path.dirname(os.path.abspath(args.out))
    os.makedirs(out_dir, exist_ok=True)
    prs.save(args.out)

    # 封装 PDF(可选)
    if args.pdf:
        imgs = []
        for p in pages:
            with Image.open(p) as im:
                imgs.append(im.convert("RGB"))
        first, rest = imgs[0], imgs[1:]
        pdf_dir = os.path.dirname(os.path.abspath(args.pdf))
        os.makedirs(pdf_dir, exist_ok=True)
        first.save(args.pdf, save_all=True, append_images=rest)

    # 交付验证:向后验证(重开 pptx 数页数、数图)
    check = Presentation(args.out)
    n_pic = sum(1 for s in check.slides for sh in s.shapes if sh.shape_type == 13)
    print(f"OK {args.out} · {len(check.slides)} 页(含 {n_pic} 张整图)"
          + (f" · {args.pdf}" if args.pdf else ""))
    if args.pdf and os.path.exists(args.pdf):
        print(f"OK PDF {args.pdf} · {os.path.getsize(args.pdf) // 1024} KB")
    print("NOT RUN: 交付前抽 1–2 页人工/visual-judge 复核生图文字是否准确"
          "(生图模型文字易错,这是图片版路线的头号质检点)")


if __name__ == "__main__":
    main()
